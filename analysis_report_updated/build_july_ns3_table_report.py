#!/usr/bin/env python3
"""July NS-3 progress presentation — mixed layout with logo space."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

OUT = Path(__file__).resolve().parent / "July_NS3_TrafficQoS_Progress.pptx"

NAVY = RGBColor(0x1F, 0x4E, 0x79)
DARK = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x5F, 0x6B, 0x7A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SKY = RGBColor(0xE9, 0xF2, 0xF9)
GREEN = RGBColor(0x16, 0x65, 0x34)
LGREEN = RGBColor(0xF0, 0xFD, 0xF4)
RED = RGBColor(0xB9, 0x1C, 0x1C)
LRED = RGBColor(0xFD, 0xF2, 0xF2)
AMBER = RGBColor(0xB4, 0x53, 0x09)
LAMBER = RGBColor(0xFF, 0xF8, 0xE7)
LINE = RGBColor(0xCF, 0xD8, 0xE3)
FONT = "Arial"

AUTHOR = "Sheikh Sayed Bin Rahman"
LAB = "PIC Lab, Kyushu Institute of Technology (KIT)"
PROJECT = "Hybrid Network Construction and Performance Evaluation\nfor Multi-Robot Systems at Construction Sites"


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def rect(slide, l, t, w, h, fill=None, line=None, rounded=False):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE, l, t, w, h
    )
    if rounded:
        try:
            shape.adjustments[0] = 0.08
        except Exception:
            pass
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(1)
    shape.shadow.inherit = False
    return shape


def textbox(slide, text, l, t, w, h, size=14, bold=False, color=DARK,
            align=PP_ALIGN.LEFT, italic=False):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Pt(2)
    tf.margin_right = Pt(2)
    tf.margin_top = Pt(0)
    tf.margin_bottom = Pt(0)
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    return box


def logo_slot(slide, cover=False):
    """Empty top-left box reserved for lab logo."""
    size = Inches(1.05) if cover else Inches(0.78)
    left = Inches(0.35)
    top = Inches(0.25)
    box = rect(slide, left, top, size, size, fill=WHITE, line=LINE, rounded=True)
    textbox(
        slide,
        "LAB\nLOGO",
        left,
        top + Inches(0.22 if cover else 0.15),
        size,
        Inches(0.6),
        size=10 if cover else 8,
        bold=True,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )
    return box


def page_header(slide, title, subtitle=None):
    logo_slot(slide)
    textbox(slide, title, Inches(1.35), Inches(0.28), Inches(8.3), Inches(0.45),
            size=20, bold=True, color=NAVY)
    if subtitle:
        textbox(slide, subtitle, Inches(1.35), Inches(0.72), Inches(8.3), Inches(0.3),
                size=11, color=MUTED)
    rect(slide, Inches(1.35), Inches(1.08), Inches(1.0), Pt(3), fill=AMBER)


def footer(slide, page, total=8):
    textbox(
        slide,
        f"{AUTHOR}  ·  {LAB}",
        Inches(0.4),
        Inches(7.15),
        Inches(7.2),
        Inches(0.22),
        size=8,
        color=MUTED,
    )
    textbox(
        slide,
        f"{page} / {total}",
        Inches(8.6),
        Inches(7.15),
        Inches(1.0),
        Inches(0.22),
        size=8,
        color=MUTED,
        align=PP_ALIGN.RIGHT,
    )


def set_cell(cell, value, fill, color=DARK, bold=False, size=10, align=PP_ALIGN.CENTER):
    cell.text = str(value)
    cell.fill.solid()
    cell.fill.fore_color.rgb = fill
    cell.margin_left = Pt(4)
    cell.margin_right = Pt(4)
    cell.margin_top = Pt(3)
    cell.margin_bottom = Pt(3)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    for p in cell.text_frame.paragraphs:
        p.alignment = align
        for r in p.runs:
            r.font.name = FONT
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = color


def add_table(slide, rows, left, top, width, height, widths=None, font_size=10):
    shape = slide.shapes.add_table(len(rows), len(rows[0]), left, top, width, height)
    table = shape.table
    if widths:
        for i, w in enumerate(widths):
            table.columns[i].width = w
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            if ri == 0:
                fill, color, bold = NAVY, WHITE, True
            else:
                fill = SKY if ri % 2 == 0 else WHITE
                color, bold = DARK, ci == 0
            align = PP_ALIGN.LEFT if ci == 0 else PP_ALIGN.CENTER
            set_cell(table.cell(ri, ci), val, fill, color, bold, font_size, align)
    return shape


def card(slide, l, t, w, h, title, body_lines, accent=NAVY):
    rect(slide, l, t, w, h, fill=WHITE, line=LINE, rounded=True)
    rect(slide, l, t, w, Inches(0.08), fill=accent)
    textbox(slide, title, l + Inches(0.18), t + Inches(0.2), w - Inches(0.3), Inches(0.35),
            size=13, bold=True, color=NAVY)
    y = t + Inches(0.55)
    for line in body_lines:
        textbox(slide, "•  " + line, l + Inches(0.18), y, w - Inches(0.3), Inches(0.32),
                size=11, color=DARK)
        y += Inches(0.32)


def metric(slide, l, t, w, h, value, label, accent=NAVY):
    rect(slide, l, t, w, h, fill=WHITE, line=LINE, rounded=True)
    rect(slide, l, t, Pt(5), h, fill=accent)
    textbox(slide, value, l + Inches(0.15), t + Inches(0.18), w - Inches(0.25), Inches(0.5),
            size=24, bold=True, color=DARK, align=PP_ALIGN.CENTER)
    textbox(slide, label, l + Inches(0.12), t + Inches(0.72), w - Inches(0.2), Inches(0.45),
            size=10, color=MUTED, align=PP_ALIGN.CENTER)


# ── Slides ───────────────────────────────────────────────────────────────────

def slide_cover(prs):
    s = blank(prs)
    rect(s, Inches(0), Inches(0), Inches(10), Inches(7.5), fill=NAVY)
    logo_slot(s, cover=True)

    textbox(s, "JULY 2026 PROGRESS UPDATE", Inches(1.6), Inches(1.55), Inches(7.5), Inches(0.35),
            size=12, bold=True, color=AMBER, align=PP_ALIGN.CENTER)
    textbox(s, "NS-3 Traffic QoS Model &\nHybrid Switching Reliability",
            Inches(0.8), Inches(2.1), Inches(8.4), Inches(1.15),
            size=28, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    textbox(s, PROJECT, Inches(1.2), Inches(3.4), Inches(7.6), Inches(0.7),
            size=13, color=SKY, align=PP_ALIGN.CENTER)

    # Name card
    rect(s, Inches(2.0), Inches(4.5), Inches(6.0), Inches(1.55), fill=WHITE, rounded=True)
    textbox(s, "Presented by", Inches(2.2), Inches(4.65), Inches(5.6), Inches(0.25),
            size=10, color=MUTED, align=PP_ALIGN.CENTER)
    textbox(s, AUTHOR, Inches(2.2), Inches(4.95), Inches(5.6), Inches(0.35),
            size=16, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    textbox(s, LAB, Inches(2.2), Inches(5.35), Inches(5.6), Inches(0.3),
            size=11, color=DARK, align=PP_ALIGN.CENTER)
    textbox(s, "NS-3.45  ·  WiFi Mesh + LTE / 5G NR", Inches(2.2), Inches(5.7), Inches(5.6), Inches(0.25),
            size=10, color=MUTED, align=PP_ALIGN.CENTER)

    textbox(s, "Replace the top-left box with the PIC Lab logo",
            Inches(0.4), Inches(6.95), Inches(9.2), Inches(0.25),
            size=9, italic=True, color=SKY, align=PP_ALIGN.CENTER)


def slide_overview(prs):
    s = blank(prs)
    page_header(s, "What was completed in July",
                "Software implementation and reliability fixes for the NS-3 hybrid simulator")

    items = [
        ("1", "QoS traffic model",
         "Replaced generic traffic with Control, Sensor, and Video flows, each with DSCP marking."),
        ("2", "Per-flow measurement",
         "Added flow_metrics.py for Control PDR, P99 latency, and throughput share."),
        ("3", "WiFi association crash fix",
         "Patched stale Association Response handling so multi-AP roaming no longer aborts NS-3."),
        ("4", "Switch recovery improvement",
         "Recovery confirmation now uses Control UDP, improving resolved-switch rate on test cases."),
        ("5", "Multi-seed evaluation",
         "Completed a 5G NR matrix across seeds 6/7/8 and STA counts 5/10/15."),
    ]
    y = Inches(1.35)
    for num, title, desc in items:
        rect(s, Inches(0.55), y, Inches(8.9), Inches(0.9), fill=WHITE, line=LINE, rounded=True)
        rect(s, Inches(0.55), y, Inches(0.7), Inches(0.9), fill=NAVY)
        textbox(s, num, Inches(0.55), y + Inches(0.25), Inches(0.7), Inches(0.4),
                size=18, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        textbox(s, title, Inches(1.45), y + Inches(0.15), Inches(7.7), Inches(0.3),
                size=14, bold=True, color=NAVY)
        textbox(s, desc, Inches(1.45), y + Inches(0.48), Inches(7.7), Inches(0.35),
                size=11, color=DARK)
        y += Inches(1.0)
    footer(s, 2)


def slide_qos(prs):
    s = blank(prs)
    page_header(s, "July deliverable — robot traffic QoS model",
                "Three application flows with independent FlowMonitor measurement")

    card(s, Inches(0.45), Inches(1.35), Inches(2.95), Inches(2.7), "Control", [
        "UDP, ~1 KB / 50 ms",
        "DSCP EF (robot commands)",
        "Target: P99 ≤ 50 ms",
        "Loss target: 0%",
    ], accent=RED)
    card(s, Inches(3.55), Inches(1.35), Inches(2.95), Inches(2.7), "Sensor", [
        "TCP telemetry upload",
        "DSCP AF31",
        "Target: delay ≤ 200 ms",
        "Loss target: < 5%",
    ], accent=NAVY)
    card(s, Inches(6.65), Inches(1.35), Inches(2.95), Inches(2.7), "Video", [
        "TCP continuous stream",
        "DSCP AF41",
        "Target: delay ≤ 500 ms",
        "Loss target: < 10%",
    ], accent=GREEN)

    textbox(s, "Deliverable status", Inches(0.55), Inches(4.3), Inches(8.8), Inches(0.3),
            size=13, bold=True, color=NAVY)
    rows = [
        ["Artifact", "Role", "Status"],
        ["traffic_qos.cc", "NS-3 scenario with 3-flow traffic + DSCP", "Complete"],
        ["flow_metrics.py", "Per-flow PDR / delay / P99 / throughput", "Complete"],
        ["traffic_qos_report.pdf", "Full statistical analysis report", "Pending 192 runs"],
    ]
    add_table(s, rows, Inches(0.55), Inches(4.7), Inches(8.9), Inches(1.55),
              widths=[Inches(2.3), Inches(4.6), Inches(2.0)], font_size=11)
    footer(s, 3)


def slide_crash(prs):
    s = blank(prs)
    page_header(s, "WiFi BSSID crash — root cause and fix",
                "NS-3 aborted under multi-AP roaming with shared SSID MeshHotspot")

    # Timeline style cards
    steps = [
        ("1", "Associate", "STA requests association to AP 0c"),
        ("2", "Roam", "Missed beacons → rescan → new request to AP 0b"),
        ("3", "Stale frame", "Delayed AssocResp from abandoned AP 0c arrives"),
        ("4", "Crash", "Assert: target BSSID ≠ response BSSID"),
    ]
    x = Inches(0.45)
    for num, title, desc in steps:
        rect(s, x, Inches(1.35), Inches(2.2), Inches(1.85), fill=WHITE, line=LINE, rounded=True)
        textbox(s, num, x + Inches(0.15), Inches(1.5), Inches(0.4), Inches(0.35),
                size=16, bold=True, color=NAVY)
        textbox(s, title, x + Inches(0.15), Inches(1.95), Inches(1.9), Inches(0.3),
                size=13, bold=True, color=DARK)
        textbox(s, desc, x + Inches(0.15), Inches(2.35), Inches(1.9), Inches(0.65),
                size=11, color=MUTED)
        x += Inches(2.35)

    rect(s, Inches(0.55), Inches(3.5), Inches(8.9), Inches(1.15), fill=LGREEN, line=LINE, rounded=True)
    textbox(s, "Correction in sta-wifi-mac.cc", Inches(0.75), Inches(3.65), Inches(8.5), Inches(0.3),
            size=13, bold=True, color=GREEN)
    textbox(s,
            "Ignore Association Responses whose BSSID does not match the STA’s current target AP. "
            "Keep waiting for the correct response instead of aborting the simulation.",
            Inches(0.75), Inches(4.05), Inches(8.5), Inches(0.45),
            size=12, color=DARK)

    metric(s, Inches(0.55), Inches(4.95), Inches(2.85), Inches(1.4), "4 / 9", "runs crashed before", RED)
    metric(s, Inches(3.55), Inches(4.95), Inches(2.85), Inches(1.4), "9 / 9", "runs complete after", GREEN)
    metric(s, Inches(6.55), Inches(4.95), Inches(2.85), Inches(1.4), "0", "aborts in re-sweep", NAVY)
    footer(s, 4)


def slide_recovery(prs):
    s = blank(prs)
    page_header(s, "Switch recovery confirmation improved",
                "Representative case: 5G NR · 10 STA · seed 7 · 60 seconds")

    textbox(s, "What changed", Inches(0.55), Inches(1.3), Inches(8.8), Inches(0.3),
            size=13, bold=True, color=NAVY)
    textbox(s,
            "Recovery was previously confirmed only when Sensor TCP traffic resumed. "
            "TCP can stall after a route rewrite, so many successful handovers were counted as timeouts. "
            "Recovery is now confirmed from the continuous Control UDP flow, and stale pending switch "
            "events are retired when a newer path decision occurs.",
            Inches(0.55), Inches(1.65), Inches(8.9), Inches(0.9),
            size=12, color=DARK)

    # Before / After
    rect(s, Inches(0.45), Inches(2.75), Inches(4.4), Inches(2.55), fill=LRED, line=LINE, rounded=True)
    textbox(s, "Before", Inches(0.65), Inches(2.95), Inches(4.0), Inches(0.3),
            size=14, bold=True, color=RED)
    textbox(s, "57%", Inches(0.65), Inches(3.35), Inches(4.0), Inches(0.55),
            size=32, bold=True, color=DARK)
    textbox(s, "resolved switches\n20 timeouts out of 47 events\nmedian interruption 178 ms",
            Inches(0.65), Inches(4.0), Inches(4.0), Inches(1.0),
            size=12, color=DARK)

    rect(s, Inches(5.15), Inches(2.75), Inches(4.4), Inches(2.55), fill=LGREEN, line=LINE, rounded=True)
    textbox(s, "After", Inches(5.35), Inches(2.95), Inches(4.0), Inches(0.3),
            size=14, bold=True, color=GREEN)
    textbox(s, "87%", Inches(5.35), Inches(3.35), Inches(4.0), Inches(0.55),
            size=32, bold=True, color=DARK)
    textbox(s, "resolved switches\n2 timeouts out of 47 events\nmedian interruption 55 ms",
            Inches(5.35), Inches(4.0), Inches(4.0), Inches(1.0),
            size=12, color=DARK)

    textbox(s,
            "Full 9-run matrix (all seeds/STA): 73% overall resolved · WiFi→NR 88% · NR→WiFi return 57%",
            Inches(0.55), Inches(5.55), Inches(8.9), Inches(0.45),
            size=12, color=MUTED)
    textbox(s,
            "Note: return from cellular to WiFi is still harder than the forward WiFi→NR path; "
            "multi-seed data is required for the report.",
            Inches(0.55), Inches(6.05), Inches(8.9), Inches(0.45),
            size=11, color=DARK)
    footer(s, 5)


def slide_seeds(prs):
    s = blank(prs)
    page_header(s, "Seeds and STA count still matter",
                "Same code, same 60 s / 5G NR setup — different RNG seeds change outcomes")

    rows = [
        ["Seed", "STA", "Switches", "Resolved", "Timeout", "Timeout %", "Median ms"],
        ["6", "5", "24", "14", "6", "25%", "65"],
        ["6", "10", "60", "42", "8", "13%", "61"],
        ["6", "15", "96", "72", "16", "17%", "245"],
        ["7", "5", "23", "15", "4", "17%", "564"],
        ["7", "10", "47", "41", "2", "4%", "55"],
        ["7", "15", "84", "62", "13", "15%", "63"],
        ["8", "5", "34", "29", "2", "6%", "39"],
        ["8", "10", "60", "50", "6", "10%", "52"],
        ["8", "15", "78", "45", "18", "23%", "557"],
    ]
    add_table(s, rows, Inches(0.4), Inches(1.3), Inches(9.2), Inches(4.35),
              widths=[Inches(0.85), Inches(0.85), Inches(1.3), Inches(1.3),
                      Inches(1.2), Inches(1.35), Inches(1.35)],
              font_size=10)

    textbox(s,
            "Best: seed 7 / 10 STA (4% timeout). Worst: seed 6 / 5 STA (25%). "
            "A single seed would have overstated reliability — hence the 192-run campaign.",
            Inches(0.5), Inches(5.9), Inches(9.0), Inches(0.55),
            size=12, color=DARK)
    footer(s, 6)


def slide_192(prs):
    s = blank(prs)
    page_header(s, "Next: 192-run campaign for the formal report",
                "Phase 1 factorial design with the July QoS traffic model")

    textbox(s,
            "The remaining July analysis deliverable is traffic_qos_report.pdf. "
            "It needs a full multi-factor campaign, not single-run results.",
            Inches(0.55), Inches(1.3), Inches(8.9), Inches(0.5),
            size=12, color=DARK)

    rows = [
        ["Factor", "Values", "Levels"],
        ["Cellular mode", "LTE, 5G NR", "2"],
        ["WiFi hotspot band", "2.4 GHz, 5 GHz", "2"],
        ["STA count", "5, 10, 15, 20", "4"],
        ["Payload / load", "10 KB, 50 KB, 1 MB, 2 MB", "4"],
        ["RNG seed", "6, 7, 8", "3"],
        ["Total", "2 × 2 × 4 × 4 × 3", "192 runs"],
    ]
    add_table(s, rows, Inches(0.7), Inches(1.95), Inches(8.6), Inches(2.85),
              widths=[Inches(2.5), Inches(4.4), Inches(1.7)], font_size=11)

    card(s, Inches(0.45), Inches(5.1), Inches(4.4), Inches(1.5), "Campaign outputs", [
        "Control PDR and P99 tables",
        "LTE vs NR comparison",
        "Switch interruption statistics",
    ], accent=NAVY)
    card(s, Inches(5.15), Inches(5.1), Inches(4.4), Inches(1.5), "Report goal", [
        "Seed-averaged KPIs",
        "Factor-level confidence",
        "traffic_qos_report.pdf",
    ], accent=GREEN)
    footer(s, 7)


def slide_thanks(prs):
    s = blank(prs)
    rect(s, Inches(0), Inches(0), Inches(10), Inches(7.5), fill=NAVY)
    logo_slot(s, cover=True)

    textbox(s, "Thank You", Inches(0.8), Inches(2.4), Inches(8.4), Inches(0.7),
            size=40, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    textbox(s, "Questions and discussion are welcome.",
            Inches(1.2), Inches(3.25), Inches(7.6), Inches(0.4),
            size=16, color=SKY, align=PP_ALIGN.CENTER)

    rect(s, Inches(2.2), Inches(4.2), Inches(5.6), Inches(1.55), fill=WHITE, rounded=True)
    textbox(s, AUTHOR, Inches(2.4), Inches(4.4), Inches(5.2), Inches(0.35),
            size=15, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    textbox(s, LAB, Inches(2.4), Inches(4.8), Inches(5.2), Inches(0.3),
            size=11, color=DARK, align=PP_ALIGN.CENTER)
    textbox(s, "Hybrid Robot Network Simulator  ·  NS-3.45",
            Inches(2.4), Inches(5.2), Inches(5.2), Inches(0.3),
            size=11, color=MUTED, align=PP_ALIGN.CENTER)

    textbox(s, "Replace the top-left box with the PIC Lab logo",
            Inches(0.4), Inches(6.95), Inches(9.2), Inches(0.25),
            size=9, italic=True, color=SKY, align=PP_ALIGN.CENTER)


def main():
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    slide_cover(prs)
    slide_overview(prs)
    slide_qos(prs)
    slide_crash(prs)
    slide_recovery(prs)
    slide_seeds(prs)
    slide_192(prs)
    slide_thanks(prs)

    prs.save(OUT)
    print(f"Wrote {OUT}")
    print(f"Slides: {len(prs.slides)}")
    print(f"Size: {OUT.stat().st_size / 1024:.1f} KB")


if __name__ == "__main__":
    main()
